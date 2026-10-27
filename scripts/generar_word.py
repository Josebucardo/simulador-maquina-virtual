"""Genera documentos Word editables desde las fuentes Markdown de la entrega."""
from pathlib import Path
import re
import sys
import textwrap

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import generar_documento as pdf_builder


def rich(paragraph, text):
    tokens = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        bold = token.startswith('**') and token.endswith('**')
        italic = not bold and token.startswith('*') and token.endswith('*')
        code = token.startswith('`') and token.endswith('`')
        token = token[2:-2] if bold else token[1:-1] if italic or code else token
        token = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', token)
        run = paragraph.add_run(token)
        run.bold, run.italic = bold, italic
        if code:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
    return paragraph


def configure(document):
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    # Suprimir reglas y colores heredados del estilo de Office local.
    for style in document.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
    for name, size in [('Title', 25), ('Subtitle', 13), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 11)]:
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = 'Arial', Pt(size), RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    document.styles['Caption'].font.color.rgb = RGBColor(60, 60, 60)
    code = document.styles.add_style('Código ASM', 1)
    code.font.name, code.font.size = 'Consolas', Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.keep_with_next = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Grupo 13 · Página ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.author = 'Grupo 13'
    document.core_properties.title = 'Máquina virtual simple'


def table(document, lines):
    rows = [pdf_builder.split_table_row(line) for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-{2,}:?', cell or ' ') for cell in row)]
    count = max(map(len, rows))
    widths = []
    for col in range(count):
        values = [len(row[col]) if col < len(row) else 0 for row in rows]
        widths.append(max(8, min(40, max(len(rows[0][col]), sorted(values)[int((len(values)-1)*.8)]))))
    widths = [17.4 * value / sum(widths) for value in widths]
    for col, heading in enumerate(rows[0]):
        if heading.strip().lower() == 'estado' and widths[col] < 2.3:
            deficit = 2.3 - widths[col]
            total_others = sum(widths) - widths[col]
            widths = [2.3 if index == col else width - deficit * width / total_others
                      for index, width in enumerate(widths)]
    result = document.add_table(rows=0, cols=count)
    result.autofit = False
    for index, row in enumerate(rows):
        cells = result.add_row().cells
        for col, cell in enumerate(cells):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    edge.set(qn('w:' + key), value)
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:w'), '75')
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            props.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '18384E' if index == 0 else 'F2F5F7' if index % 2 == 0 else 'FFFFFF')
            props.append(shading)
            p = rich(cell.paragraphs[0], row[col] if col < len(row) else '')
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.size = Pt(8.5 if count < 5 else 8)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if index:
                no_split = OxmlElement('w:cantSplit')
                result.rows[index]._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            header = OxmlElement('w:tblHeader')
            result.rows[index]._tr.get_or_add_trPr().append(header)
    for index, col in enumerate(result.columns):
        col.width = Cm(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def image_for_mermaid(source, number):
    """Usa los mismos diagramas del PDF, renderizados con fuentes incrustadas."""
    from reportlab.graphics import renderPDF
    import pypdfium2
    pdf_builder.register_fonts()
    drawing = pdf_builder.diagram(source)
    temp = ROOT / 'tmp/word/diagramas'
    temp.mkdir(parents=True, exist_ok=True)
    pdf = temp / f'diagrama_{number}.pdf'
    png = temp / f'diagrama_{number}.png'
    renderPDF.drawToFile(drawing, str(pdf))
    doc = pypdfium2.PdfDocument(str(pdf))
    doc[0].render(scale=3).to_pil().save(png)
    doc.close()
    return png, drawing.height / drawing.width


def markdown(document, source, base):
    lines = source.splitlines()
    index, figure = 0, 0
    number_id = None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('```'):
            language = line[3:].strip()
            end = index + 1
            while end < len(lines) and not lines[end].strip().startswith('```'):
                end += 1
            content = '\n'.join(lines[index+1:end])
            if language == 'mermaid':
                figure += 1
                path, ratio = image_for_mermaid(content, figure)
                width = min(16.5, 20.0 / ratio)
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(width))
            else:
                block = content.expandtabs(4).splitlines()
                for position, code_line in enumerate(block):
                    pieces = textwrap.wrap(code_line, width=104, subsequent_indent='    ', replace_whitespace=False, drop_whitespace=False) or ['']
                    for part in pieces:
                        p = document.add_paragraph(part, 'Código ASM')
                        if len(block) < 33 and position < len(block)-1:
                            p.paragraph_format.keep_with_next = True
                if block:
                    p.paragraph_format.space_after = Pt(6)
            index = end + 1
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading:
            level = min(3, max(1, len(heading[1])-1))
            title = re.sub(r'^Anexo ([A-D])\.\s*', r'Anexo \1 ', heading[2])
            if title.startswith('Guion del Integrante') or title.startswith('Diapositiva 9 '):
                document.add_page_break()
            document.add_heading(title, level=level)
            index += 1
            continue
        if line.startswith('|') and index+1 < len(lines) and re.match(r'^\s*\|?\s*:?-{2,}', lines[index+1]):
            end = index+1
            while end < len(lines) and lines[end].strip().startswith('|'):
                end += 1
            table(document, lines[index:end])
            index = end
            continue
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if image:
            p = document.add_paragraph()
            p.add_run().add_picture(str((base/image[2]).resolve()), width=Cm(16.8))
            document.add_paragraph(image[1], 'Caption')
            index += 1
            continue
        item = re.match(r'^(?:[-*]|\d+[.)])\s+(.+)$', line)
        if item:
            if line[0].isdigit():
                start = int(re.match(r'\d+', line)[0])
                if number_id is None or start == 1:
                    numbering = document.part.numbering_part.element
                    existing = [int(v.get(qn('w:abstractNumId'))) for v in numbering.findall(qn('w:abstractNum'))]
                    abstract_id = max(existing, default=-1) + 1
                    abstract = OxmlElement('w:abstractNum')
                    abstract.set(qn('w:abstractNumId'), str(abstract_id))
                    level = OxmlElement('w:lvl')
                    level.set(qn('w:ilvl'), '0')
                    for tag, value in [('start', str(start)), ('numFmt', 'decimal'), ('lvlText', '%1.')]:
                        child = OxmlElement('w:' + tag)
                        child.set(qn('w:val'), value)
                        level.append(child)
                    abstract.append(level)
                    first_num = numbering.find(qn('w:num'))
                    numbering.insert(numbering.index(first_num) if first_num is not None else len(numbering), abstract)
                    number_id = numbering.add_num(abstract_id).numId
                p = document.add_paragraph(style='List Paragraph')
                properties = p._p.get_or_add_pPr()
                number = OxmlElement('w:numPr')
                for tag, value in [('ilvl', '0'), ('numId', str(number_id))]:
                    child = OxmlElement('w:' + tag)
                    child.set(qn('w:val'), value)
                    number.append(child)
                properties.append(number)
                p.paragraph_format.left_indent = Cm(.65)
                p.paragraph_format.first_line_indent = Cm(-.45)
            else:
                p = document.add_paragraph(style='List Bullet')
            rich(p, item[1])
            index += 1
            continue
        group = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r'^\s*(?:#|```|\||[-*]\s|\d+[.)]\s|!\[)', lines[index]):
            group.append(lines[index].strip())
            index += 1
        rich(document.add_paragraph(), ' '.join(group))


def technical():
    doc = Document()
    configure(doc)
    source = (ROOT/'docs/DOCUMENTO_TECNICO_GRUPO13.md').read_text(encoding='utf-8')
    _, cover, body = pdf_builder.extract_cover(source)
    doc.add_heading('Implementación de una máquina virtual simple', 0)
    doc.add_paragraph('Documento técnico y anexos', 'Subtitle')
    for line in cover.splitlines():
        if line.strip():
            rich(doc.add_paragraph(), line.strip())
    doc.add_page_break()
    markdown(doc, body, ROOT/'docs')
    for name, label in [('MANUAL_USUARIO.md', 'Anexo A Manual de usuario'), ('MANUAL_TECNICO.md', 'Anexo B Manual técnico'), ('EVIDENCIAS_PRUEBAS.md', 'Anexo C Evidencias de pruebas'), ('GUIA_DEFENSA.md', 'Anexo D Guía de defensa')]:
        doc.add_page_break()
        doc.add_heading(label, 1)
        content = (ROOT/'docs'/name).read_text(encoding='utf-8')
        content = re.sub(r'^# .+\n?', '', content, count=1)
        start_paragraph = len(doc.paragraphs)
        markdown(doc, content, ROOT/'docs')
        if name == 'MANUAL_USUARIO.md':
            for paragraph in doc.paragraphs[start_paragraph:]:
                if paragraph.style.name != 'Código ASM':
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.line_spacing = 1.08
                if paragraph.style.name.startswith('Heading'):
                    paragraph.paragraph_format.space_before = Pt(10)
    out = ROOT/'output/word/ENTREGA_GRUPO13.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


def exposition():
    source_path = ROOT/'docs/exposicion/GUIA_TRES_INTEGRANTES.md'
    if not source_path.exists():
        return
    doc = Document()
    configure(doc)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(5)
    source = source_path.read_text(encoding='utf-8')
    title = re.search(r'^# (.+)', source)
    doc.add_heading(title[1] if title else 'Guía de exposición para tres integrantes', 0)
    source = re.sub(r'^# .+\n?', '', source, count=1)
    markdown(doc, source, source_path.parent)
    out = ROOT/'output/exposicion/GUIA_TRES_INTEGRANTES.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


def user_manual():
    """Documento independiente para abrir el programa y usar sus controles."""
    source_path = ROOT / 'docs/MANUAL_USUARIO.md'
    doc = Document()
    configure(doc)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(10)
    source = source_path.read_text(encoding='utf-8')
    doc.add_heading('Manual de usuario de VM Studio', 0)
    doc.add_paragraph('Máquina virtual simple del grupo 13', 'Subtitle')
    source = re.sub(r'^# .+\n?', '', source, count=1)
    markdown(doc, source, source_path.parent)
    out = ROOT / 'output/manual/MANUAL_USUARIO_VM_STUDIO.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == '__main__':
    technical()
    exposition()
    user_manual()
